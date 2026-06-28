"""Generate KB FinAgent Culture tech stack presentation."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from generate_workflow_ppt import (
    add_content_slide,
    add_flow_slide,
    add_section_slide,
    add_table_slide,
    add_title_slide,
)

OUTPUT = Path(__file__).parent / "KB_FinAgent_Culture_주요기술.pptx"


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "KB FinAgent (Culture)",
        "주요 기술 스택 소개",
    )

    add_content_slide(
        prs,
        "한 줄 요약",
        [
            "Flask 웹 앱 + PostgreSQL(Aurora) 데이터 조회",
            "+ LangGraph 질문 흐름 제어",
            "+ AWS Bedrock(Claude) AI 응답",
            "+ Vanilla JS / Chart.js 채팅 UI",
        ],
        font_size=20,
    )

    add_flow_slide(
        prs,
        "전체 아키텍처",
        [
            "[브라우저]  HTML + culture_chat.js + Chart.js",
            "      ↕  REST API / SSE 스트리밍",
            "[Flask]  culture_app.py  (로그인, 세션, API)",
            "      ↕",
            "[LangGraph]  culture_workflow.py  (질문 분기)",
            "      ↕",
            "[culture_inst1_agents.py]  SQL 생성, 에이전트 로직",
            "      ↕                    ↕",
            "[PostgreSQL]            [AWS Bedrock / Claude]",
            " Aurora INST1 테이블      질문 분석·요약·답변",
        ],
    )

    add_section_slide(prs, "백엔드")

    add_table_slide(
        prs,
        "백엔드 — Python / Flask",
        ["기술", "역할"],
        [
            ("Python 3.12", "전체 서버 로직"),
            ("Flask", "웹 서버, 로그인·채팅 API, HTML 화면"),
            ("Werkzeug", "세션, 비밀번호 해시, 프록시(ProxyFix)"),
            ("python-dotenv", ".env.local 환경 변수 로드"),
        ],
        col_widths=[2.8, 6.2],
    )

    add_content_slide(
        prs,
        "백엔드 실행 방식",
        [
            "로컬: culture_app.py 직접 실행 → http://127.0.0.1:5051",
            "프로덕션: Vercel serverless (api/index.py)",
            "run_culture.bat 으로 Windows에서 간편 실행 가능",
        ],
        font_size=18,
    )

    add_section_slide(prs, "AI · 워크플로우")

    add_table_slide(
        prs,
        "AI · 워크플로우",
        ["기술", "역할"],
        [
            ("LangGraph", "질문 분석 → 기능별 분기 → 응답 조립 (그래프)"),
            ("AWS Bedrock", "Claude Sonnet 4.5 등 LLM 호출"),
            ("boto3", "Bedrock API 클라이언트, 리전·모델 폴백"),
        ],
        col_widths=[2.5, 6.5],
    )

    add_content_slide(
        prs,
        "LangGraph + Bedrock 활용",
        [
            "질문 종류(컬럼 설명, 집계, 차트 등)에 따라 노드가 분기됩니다.",
            "각 노드에서 Bedrock을 호출해 질문 분류, SQL 생성, 데이터 요약을 수행합니다.",
            "기본 모델: anthropic.claude-sonnet-4-5-20250929-v1:0",
            "환경 변수 BEDROCK_MODEL_ID 로 모델 변경 가능",
        ],
        font_size=17,
    )

    add_section_slide(prs, "데이터베이스")

    add_table_slide(
        prs,
        "데이터베이스",
        ["기술", "역할"],
        [
            ("PostgreSQL", "데이터 저장 (Aurora 호스팅)"),
            ("psycopg2", "Python에서 SQL 실행·DB 연결"),
            ("INST1 스키마", "그룹고객 분석 테이블 3종"),
            ("public.회원", "로그인용 회원 테이블"),
        ],
        col_widths=[2.5, 6.5],
    )

    add_table_slide(
        prs,
        "INST1 분석 테이블",
        ["테이블", "내용"],
        [
            ("TSHDEOA01", "그룹고객기본정보"),
            ("TSHDEOA02", "그룹고객거래기본"),
            ("TSHDE0ZCD", "그룹고객분석인스턴스목록"),
        ],
        col_widths=[2.8, 6.2],
    )

    add_section_slide(prs, "프론트엔드")

    add_table_slide(
        prs,
        "프론트엔드",
        ["기술", "역할"],
        [
            ("HTML / CSS", "Flask가 렌더링하는 채팅 UI (React/Vue 미사용)"),
            ("Vanilla JavaScript", "culture_chat.js — 메시지·표·추천 질문"),
            ("Chart.js (CDN)", "집계 결과 막대그래프"),
            ("SSE", "/api/chat/stream — 스트리밍 응답"),
        ],
        col_widths=[2.8, 6.2],
    )

    add_section_slide(prs, "클라우드 · 배포")

    add_table_slide(
        prs,
        "클라우드 · 배포",
        ["기술", "역할"],
        [
            ("Vercel", "프로덕션 배포 (vercel.json, Python serverless)"),
            ("AWS IAM", "Bedrock 호출 권한 관리"),
            ("AWS S3 (선택)", "요약 PDF 업로드·다운로드 링크"),
        ],
        col_widths=[2.5, 6.5],
    )

    add_content_slide(
        prs,
        "주요 환경 변수",
        [
            "AURORA_DB_URL — Postgres 연결 문자열",
            "FLASK_SECRET_KEY — 세션 암호화 키",
            "AWS_ACCESS_KEY_ID / AWS_SECRET_ACCESS_KEY — Bedrock 인증",
            "AWS_REGION / AWS_BEDROCK_REGION — AWS 리전",
            "BEDROCK_MODEL_ID — (선택) LLM 모델 지정",
            "CULTURE_S3_BUCKET — (선택) PDF 저장용 S3",
        ],
        font_size=16,
    )

    add_section_slide(prs, "부가 라이브러리")

    add_table_slide(
        prs,
        "부가 기능 라이브러리",
        ["기술", "역할"],
        [
            ("fpdf2", "요약 내용 PDF 생성"),
            ("matplotlib", "PDF 내 차트 이미지 렌더링"),
            ("Noto Sans KR / Nanum Gothic", "PDF 한글 폰트"),
        ],
        col_widths=[3.2, 5.8],
    )

    add_table_slide(
        prs,
        "requirements.txt 핵심 패키지",
        ["패키지", "용도"],
        [
            ("Flask", "웹 프레임워크"),
            ("langgraph", "AI 워크플로우"),
            ("boto3", "AWS Bedrock 연동"),
            ("psycopg2-binary", "PostgreSQL 연결"),
            ("fpdf2 / matplotlib", "PDF 생성 (선택)"),
        ],
        col_widths=[3.0, 6.0],
    )

    add_content_slide(
        prs,
        "핵심 파일 구조",
        [
            "culture_app.py — Flask 앱 (채팅 UI, API, SSE)",
            "culture_workflow.py — LangGraph 노드·라우팅",
            "culture_inst1_agents.py — 질문 분석, SQL, 에이전트 로직",
            "static/culture_chat.js — 채팅 UI, 차트 렌더링",
            "culture_db/ — DDL, 시드, 테이블 설정, 회원 인증",
        ],
        font_size=17,
    )

    add_content_slide(
        prs,
        "정리",
        [
            "전통적인 Flask 웹앱 위에 LangGraph로 AI 에이전트 흐름을 설계",
            "Bedrock LLM + Postgres 데이터를 연결한 그룹고객 분석 챗봇",
            "프론트는 단순(Vanilla JS), 복잡한 로직은 백엔드·워크플로우에 집중",
        ],
        font_size=19,
    )

    add_title_slide(prs, "감사합니다", "KB FinAgent Culture")

    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
